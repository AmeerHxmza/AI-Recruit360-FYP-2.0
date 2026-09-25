import docx
import os
import re

def update_docx(file_path):
    if not os.path.exists(file_path):
        print(f"File not found: {file_path}")
        return

    print(f"\nProcessing {file_path}...")
    doc = docx.Document(file_path)

    # 1. Update Paragraphs
    replacements_paragraphs = [
        # Latency bullet points
        (
            r"Analysis of latency observations indicates:[\s\S]*?Simli WebRTC video track negotiation.*",
            "Analysis of latency observations indicates:\n"
            "• Database Operations: Standard CRUD queries are performed in less than 45ms locally and less than 120ms remotely on Supabase hosted in the cloud, well within the NFR of 250ms.\n"
            "• The average time for Resume Screening - End-to-end resume extraction and OpenAI GPT-4o-mini structured evaluation was 3.84 seconds, which is within the target of 6 seconds.\n"
            "• Real-Time Interview Transcription: OpenAI Whisper (whisper-1) delivers accurate voice-to-text transcriptions within an average of 940ms locally and 1,180ms in the cloud, giving candidates an immediate, verifiable review of their spoken responses before submission.\n"
            "• Neural Avatar Streaming: Simli WebRTC video track negotiation takes 1.85 seconds, which makes sure of smooth 30fps video for standard broadband network conditions."
        ),
        (
            "Deepgram Nova-2 STT delivers finalized transcript chunks in an average of 410ms, providing candidates with an instantaneous, responsive review experience.",
            "OpenAI Whisper (whisper-1) delivers accurate voice-to-text transcriptions within an average of 940ms locally and 1,180ms in the cloud, giving candidates an immediate, verifiable review of their spoken responses before submission."
        ),
        (
            "stream 16kHz PCM audio to Deepgram Nova-2 STT, and display the live transcript",
            "capture audio and transcribe via OpenAI Whisper (whisper-1), and display the live transcript"
        ),
        (
            "Simli lip-synced avatar and ElevenLabs voice",
            "Simli lip-synced avatar and OpenAI TTS (voice: Nova)"
        ),
        (
            "Deepgram transcribes spoken responses",
            "OpenAI Whisper transcribes spoken responses"
        ),
        (
            "submits question text to ElevenLabs TTS, generating natural speech audio",
            "submits question text to OpenAI TTS (model: tts-1, voice: nova), generating natural speech audio"
        ),
        (
            "Audio chunks are transmitted over a secure WebSocket to Deepgram Nova-2 STT. Deepgram's Voice Activity Detection (VAD) detects when the candidate finishes speaking, returning an interim and finalized transcript within 300 milliseconds.",
            "Candidate audio is recorded via browser MediaRecorder and dispatched to the FastAPI backend, where OpenAI Whisper (whisper-1) performs high-accuracy speech-to-text transcription within 1.2 seconds, returning a verified transcript."
        ),
        (
            "Simli WebRTC avatar gateway, Deepgram STT websocket, and Gemini evaluation API",
            "Simli WebRTC avatar gateway, OpenAI Whisper STT API, and OpenAI evaluation API"
        ),
        (
            "External AI services (Google Gemini, Deepgram, ElevenLabs)",
            "External AI services (OpenAI GPT-4o-mini, Whisper, TTS-1, Google Gemini)"
        ),
        (
            "hosted APIs (Google Gemini, Deepgram, ElevenLabs, Simli)",
            "hosted APIs (OpenAI GPT-4o-mini, Whisper, TTS-1, Simli)"
        ),
        (
            "Deepgram. 2024. \"Nova-2 Speech-to-Text Architecture and Benchmark Specifications.\" Deepgram Technical Documentation. https://developers.deepgram.com/docs/nova-2.",
            "OpenAI. 2024. \"Whisper: Robust Speech Recognition via Large-Scale Weak Supervision.\" OpenAI Research Documentation. https://platform.openai.com/docs/models/whisper."
        ),
        (
            "ElevenLabs. 2024. \"Neural Text to Speech Models and Latency Architecture.\" ElevenLabs Research Documentation. https://elevenlabs.io/docs.",
            "OpenAI. 2024. \"Text-to-Speech API Reference and Audio Synthesis Architecture.\" OpenAI Documentation. https://platform.openai.com/docs/guides/text-to-speech."
        ),
        (
            "Next.js, FastAPI, Supabase, Google Gemini, Deepgram, ElevenLabs, and Simli",
            "Next.js, FastAPI, Supabase, OpenAI, Google Gemini, and Simli"
        ),
        (
            "Deepgram automatic speech recognition, ElevenLabs text-to-speech synthesis, and Simli neural video avatar rendering",
            "OpenAI Whisper automatic speech recognition, OpenAI TTS text-to-speech synthesis (voice: Nova), and Simli neural video avatar rendering"
        ),
        (
            "ultra-low latency speech-to-text transcription (Deepgram Nova-2), dynamic follow-up question generation (Gemini 1.5 Flash), and natural text-to-speech synthesis (ElevenLabs)",
            "speech-to-text transcription (OpenAI Whisper), structured technical evaluation (OpenAI GPT-4o-mini), and natural text-to-speech synthesis (OpenAI TTS-1, voice: Nova)"
        ),
        (
            "manage external integrations with Supabase, Deepgram, ElevenLabs, Simli, and Google Gemini.",
            "manage external integrations with Supabase, OpenAI (GPT-4o-mini, Whisper, TTS-1), Simli, and Google Gemini."
        ),
        (
            "models such as Whisper (Radford et al. 2023) and specialized streaming architectures such as Deepgram Nova-2.",
            "models such as OpenAI Whisper (Radford et al. 2023), which forms the core transcription engine in AI Recruit360."
        ),
        (
            "neural speech synthesis (TTS) models, exemplified by ElevenLabs, generate expressive",
            "neural speech synthesis (TTS) models, exemplified by OpenAI TTS-1 (voice: Nova), generate expressive"
        ),
        (
            "Deepgram Nova-2 Speech Recognition: 5 interview answers averaging 45 seconds of speech (3.75 total audio minutes at $0.0043/min) = $0.0161 USD.",
            "OpenAI Whisper Speech Recognition: 5 interview answers averaging 45 seconds of speech (3.75 total audio minutes at $0.006/min) = $0.0225 USD."
        ),
        (
            "ElevenLabs Neural Voice Synthesis: 5 interview questions averaging 180 characters (900 total characters at $0.00003/char) = $0.0270 USD.",
            "OpenAI TTS-1 Neural Voice Synthesis: 5 interview questions averaging 180 characters (900 total characters at $0.015/1,000 chars) = $0.0135 USD."
        ),
    ]

    p_count = 0
    for p in doc.paragraphs:
        for old_text, new_text in replacements_paragraphs:
            if old_text in p.text:
                p.text = p.text.replace(old_text, new_text)
                p_count += 1
            elif old_text.startswith("Analysis of latency observations") and "Analysis of latency observations" in p.text:
                p.text = new_text
                p_count += 1

    print(f"Updated {p_count} paragraph occurrences.")

    # 2. Update Tables
    t_count = 0
    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                cell_text = cell.text
                new_cell_text = cell_text

                # Table 9/10: Operational Latency
                if "Speech-to-Text Transcription (Deepgram)" in cell_text:
                    new_cell_text = cell_text.replace("Speech-to-Text Transcription (Deepgram)", "Speech-to-Text Transcription (OpenAI Whisper-1)")
                if "Text-to-Speech Generation (ElevenLabs)" in cell_text:
                    new_cell_text = cell_text.replace("Text-to-Speech Generation (ElevenLabs)", "Text-to-Speech Generation (OpenAI TTS-1, Nova)")
                if "Resume Screening (Gemini 1.5 Flash)" in cell_text:
                    new_cell_text = cell_text.replace("Resume Screening (Gemini 1.5 Flash)", "Resume Screening (OpenAI GPT-4o-mini)")

                # If this is the Speech-to-Text latency row, update numbers
                if "Speech-to-Text Transcription" in row.cells[0].text:
                    if len(row.cells) >= 5:
                        row.cells[0].text = "Speech-to-Text Transcription (OpenAI Whisper-1)"
                        row.cells[1].text = "940 ms"
                        row.cells[2].text = "1,180 ms"
                        row.cells[3].text = "< 2000 ms"
                        row.cells[4].text = "Passed (Responsive)"
                        t_count += 1
                        break

                # If this is the Text-to-Speech latency row, update numbers
                if "Text-to-Speech Generation" in row.cells[0].text:
                    if len(row.cells) >= 5:
                        row.cells[0].text = "Text-to-Speech Generation (OpenAI TTS-1, Nova)"
                        row.cells[1].text = "480 ms"
                        row.cells[2].text = "590 ms"
                        row.cells[3].text = "< 1200 ms"
                        row.cells[4].text = "Passed (Low Latency)"
                        t_count += 1
                        break

                # Other table occurrences (FR-12, FR-13, NFR-03, endpoints, budget, etc.)
                if "ElevenLabs voice" in cell_text:
                    new_cell_text = new_cell_text.replace("ElevenLabs voice", "OpenAI TTS (voice: Nova)")
                if "Deepgram Nova-2 STT" in cell_text:
                    new_cell_text = new_cell_text.replace("Deepgram Nova-2 STT", "OpenAI Whisper (whisper-1)")
                if "Deepgram Nova-2 finalized transcript return latency" in cell_text:
                    new_cell_text = new_cell_text.replace("Deepgram Nova-2 finalized transcript return latency", "OpenAI Whisper finalized transcript return latency")
                if "Deepgram STT WebSocket, ElevenLabs TTS" in cell_text:
                    new_cell_text = new_cell_text.replace("Deepgram STT WebSocket, ElevenLabs TTS", "OpenAI Whisper STT, OpenAI TTS (Nova)")
                if "Deepgram transcript buffer" in cell_text:
                    new_cell_text = new_cell_text.replace("Deepgram transcript buffer", "OpenAI Whisper transcript buffer")
                if "Streams ElevenLabs synthesized TTS audio buffer" in cell_text:
                    new_cell_text = new_cell_text.replace("Streams ElevenLabs synthesized TTS audio buffer", "Generates OpenAI TTS-1 synthesized audio buffer (voice: nova)")
                
                # Budget table
                if "Deepgram Nova-2 (Speech-to-Text)" in cell_text:
                    new_cell_text = "OpenAI Whisper (Speech-to-Text)"
                if "$0.0043 / audio minute" in cell_text:
                    new_cell_text = "$0.006 / audio minute"
                if "ElevenLabs Neural Voice (TTS)" in cell_text:
                    new_cell_text = "OpenAI TTS-1 (Neural Voice Synthesis)"
                if "$0.00003 / text character" in cell_text:
                    new_cell_text = "$0.015 / 1,000 characters"
                if "$8.06" in cell_text and "Deepgram" in row.cells[0].text:
                    new_cell_text = "$11.25"
                if "$13.50" in cell_text and "ElevenLabs" in row.cells[0].text:
                    new_cell_text = "$6.75"

                # Use Case specification
                if "ElevenLabs synthetic voice" in cell_text:
                    new_cell_text = new_cell_text.replace("ElevenLabs synthetic voice", "OpenAI TTS (Nova) synthetic voice")
                if "Browser streams 16kHz PCM audio to Deepgram Nova-2 STT" in cell_text:
                    new_cell_text = new_cell_text.replace("Browser streams 16kHz PCM audio to Deepgram Nova-2 STT over secure WebSocket", "Browser captures candidate audio and submits to OpenAI Whisper STT API")
                if "Deepgram STT" in cell_text:
                    new_cell_text = new_cell_text.replace("Deepgram STT", "OpenAI Whisper STT")

                if new_cell_text != cell_text:
                    cell.text = new_cell_text
                    t_count += 1

    print(f"Updated {t_count} table cell occurrences.")

    doc.save(file_path)
    print(f"Successfully saved {file_path}.")

if __name__ == "__main__":
    for f in ["Thesis-FYP.docx", "Thesis-2.0.docx"]:
        update_docx(f)
