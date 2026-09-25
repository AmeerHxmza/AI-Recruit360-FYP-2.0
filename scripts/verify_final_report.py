import docx

doc = docx.Document("FYP-Thesis-AIRecruit360.docx")

print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")

print("\n--- SAMPLE TABLE OF CONTENTS (Table 0) ---")
t0 = doc.tables[0]
for r in t0.rows[:12]:
    print(f"  {r.cells[0].text.strip():<45} ... {r.cells[1].text.strip()}")

print("\n--- SAMPLE LIST OF FIGURES (Table 1) ---")
t1 = doc.tables[1]
for r in t1.rows[:8]:
    print(f"  {r.cells[0].text.strip():<55} ... {r.cells[1].text.strip()}")

print("\n--- SAMPLE LIST OF TABLES (Table 2) ---")
t2 = doc.tables[2]
for r in t2.rows[:8]:
    print(f"  {r.cells[0].text.strip():<55} ... {r.cells[1].text.strip()}")

print("\n--- SAMPLE BODY PARAGRAPHS & FORMATTING ---")
for p in doc.paragraphs[70:76]:
    if p.text.strip():
        print(f"[{p.style.name} | font={p.runs[0].font.name} | size={p.runs[0].font.size.pt if p.runs and p.runs[0].font.size else 'default'} | align={p.alignment} | spacing={p.paragraph_format.line_spacing}]:\n  {p.text[:90]}...")
