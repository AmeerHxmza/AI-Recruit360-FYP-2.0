import docx

doc = docx.Document('FYP-Thesis-AIRecruit360.docx')
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt.startswith('Figure '):
        prev_txt = doc.paragraphs[i-1].text.strip() if i > 0 else ''
        has_drawing = any(r._r.xpath('.//w:drawing') for r in doc.paragraphs[i-1].runs) if i > 0 else False
        has_drawing_curr = any(r._r.xpath('.//w:drawing') for r in p.runs)
        print(f"P{i}: '{txt[:50]}' | Prev P{i-1} dw={has_drawing}: '{prev_txt[:40]}' | Curr dw={has_drawing_curr}")
