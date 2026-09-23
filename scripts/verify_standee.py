from pptx import Presentation
from pathlib import Path
import PIL.Image

ROOT = Path("d:/AI-Recruit360")
pptx_path = ROOT / "Standee-Poster-FYP.pptx"
png_path = ROOT / "Standee-Poster-FYP.png"
pdf_path = ROOT / "Standee-Poster-FYP.pdf"

prs = Presentation(str(pptx_path))
s = prs.slides[0]
print(f"Standee Dimensions: {prs.slide_width.inches:.1f} x {prs.slide_height.inches:.1f} inches")
print(f"Total Shapes: {len(s.shapes)}")

# Check image
im = PIL.Image.open(png_path)
print(f"PNG Size: {im.size} (aspect ratio = {im.size[0]/im.size[1]:.2f})")
print(f"PDF Size: {pdf_path.stat().st_size / 1024:.1f} KB")

# Create a medium-sized preview thumbnail for verification
thumb = im.resize((600, 1500), PIL.Image.Resampling.LANCZOS)
thumb.save(ROOT / ".artifacts" / "thesis" / "standee_thumbnail.png")
print("Saved thumbnail at .artifacts/thesis/standee_thumbnail.png")
