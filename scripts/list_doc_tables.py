import docx

doc = docx.Document('FYP-Thesis-AIRecruit360.docx')
print('Total tables:', len(doc.tables))
for i, t in enumerate(doc.tables):
    first_cell = t.rows[0].cells[0].text.strip().replace('\n', ' ') if t.rows else ''
    last_cell = t.rows[0].cells[-1].text.strip().replace('\n', ' ') if t.rows else ''
    print(f'Table {i}: rows={len(t.rows)}, cols={len(t.columns)}, header="{first_cell[:30]} ... {last_cell[:30]}"')
