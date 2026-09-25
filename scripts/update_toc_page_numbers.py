import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import win32com.client
import os
import re
import sys

def update_toc():
    doc_path = os.path.abspath("FYP-Thesis-AIRecruit360.docx")
    print(f"Launching Word COM on {doc_path}...", flush=True)
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0 # 0 = wdAlertsNone
    
    page_map = {}
    try:
        doc_w = word.Documents.Open(doc_path, ConfirmConversions=False, ReadOnly=False, AddToRecentFiles=False)
        pages_initial = doc_w.ComputeStatistics(2)
        print(f"Document opened. Total pages: {pages_initial}", flush=True)

        print("Harvesting page references...", flush=True)
        # Fast iteration
        for p in doc_w.Paragraphs:
            txt = p.Range.Text.strip()
            if not txt:
                continue

            # Only inspect headings, figures, tables, and front matter
            is_candidate = (
                txt.startswith("CHAPTER") or 
                txt.startswith("Chapter") or 
                txt.startswith("APPENDIX") or
                txt.startswith("Figure ") or 
                txt.startswith("Table ") or
                any(txt.startswith(fm) for fm in ["CERTIFICATION", "DEDICATION", "CONTENTS", "LIST OF FIGURES", "LIST OF TABLES", "ACKNOWLEDGMENTS", "ABSTRACT", "ABBREVIATIONS"]) or
                bool(re.match(r"^[1-7]\.[0-9]+", txt))
            )
            
            if not is_candidate:
                continue

            pg_adj = p.Range.Information(1) # wdActiveEndAdjustedPageNumber (section-adjusted!)

            # Front matter
            for fm in ["CERTIFICATION", "DEDICATION", "CONTENTS", "LIST OF FIGURES", "LIST OF TABLES", "ACKNOWLEDGMENTS", "ABSTRACT", "ABBREVIATIONS"]:
                if txt.startswith(fm) and fm not in page_map:
                    roman_numerals = ["i", "ii", "iii", "iv", "v", "vi", "vii", "viii", "ix", "x", "xi", "xii", "xiii"]
                    p_abs = p.Range.Information(3)
                    p_rom = roman_numerals[p_abs - 2] if 2 <= p_abs <= len(roman_numerals)+1 else str(p_abs)
                    page_map[fm] = p_rom

            # Chapters & Appendices
            if txt.startswith("CHAPTER") or txt.startswith("Chapter") or txt.startswith("APPENDIX"):
                ch_key = txt.split("\r")[0].strip()
                if ch_key not in page_map:
                    page_map[ch_key] = str(pg_adj)

            # Sections (1.1, 1.2...)
            m = re.match(r"^([1-7]\.[0-9]+)\s+(.+)", txt)
            if m:
                sec_code = m.group(1)
                sec_title = m.group(2).strip()
                if sec_code not in page_map:
                    page_map[sec_code] = str(pg_adj)
                    page_map[f"{sec_code} {sec_title}"] = str(pg_adj)

            # Figures
            if txt.startswith("Figure "):
                m_fig = re.match(r"^(Figure\s+[0-9A-C\.]+)", txt)
                if m_fig:
                    fig_key = m_fig.group(1)
                    if fig_key not in page_map:
                        page_map[fig_key] = str(pg_adj)

            # Tables
            if txt.startswith("Table "):
                m_tbl = re.match(r"^(Table\s+[0-9A-C\.]+)", txt)
                if m_tbl:
                    tbl_key = m_tbl.group(1)
                    if tbl_key not in page_map:
                        page_map[tbl_key] = str(pg_adj)

        doc_w.Close(False)
    finally:
        word.Quit()

    print(f"Collected {len(page_map)} page mappings successfully.", flush=True)

    # Re-inject page numbers into Table 0, 1, 2
    doc = docx.Document(doc_path)

    # Table 0: Table of Contents
    t0 = doc.tables[0]
    for row in t0.rows[1:]:
        title = row.cells[0].text.strip()
        matched_page = None
        for key, val in page_map.items():
            if key == title or title.startswith(key) or key.startswith(title):
                matched_page = val
                break
            m = re.match(r"^([1-7]\.[0-9]+)", title)
            if m and m.group(1) in page_map:
                matched_page = page_map[m.group(1)]
                break
            for ch_num in range(1, 8):
                if f"CHAPTER {ch_num}" in title and f"CHAPTER {ch_num}" in key:
                    matched_page = val
                    break
            for app_letter in ["A", "B", "C"]:
                if f"APPENDIX {app_letter}" in title and f"APPENDIX {app_letter}" in key:
                    matched_page = val
                    break
            if matched_page:
                break

        if matched_page:
            row.cells[1].text = matched_page
            p_c1 = row.cells[1].paragraphs[0]
            p_c1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for r in p_c1.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(11)

    # Table 1: List of Figures
    t1 = doc.tables[1]
    for row in t1.rows[1:]:
        title = row.cells[0].text.strip()
        matched_page = None
        for key, val in page_map.items():
            if key.startswith("Figure") and (key in title or title.startswith(key.replace("Figure ", ""))):
                matched_page = val
                break
            m = re.match(r"^([0-9A-C\.]+):", title)
            if m:
                fig_tag = f"Figure {m.group(1)}"
                if fig_tag in page_map:
                    matched_page = page_map[fig_tag]
                    break
        if matched_page:
            row.cells[1].text = matched_page
            p_c1 = row.cells[1].paragraphs[0]
            p_c1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for r in p_c1.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(11)

    # Table 2: List of Tables
    t2 = doc.tables[2]
    for row in t2.rows[1:]:
        title = row.cells[0].text.strip()
        matched_page = None
        for key, val in page_map.items():
            if key.startswith("Table") and (key in title or title.startswith(key.replace("Table ", ""))):
                matched_page = val
                break
            m = re.match(r"^([0-9A-C\.]+):", title)
            if m:
                tbl_tag = f"Table {m.group(1)}"
                if tbl_tag in page_map:
                    matched_page = page_map[tbl_tag]
                    break
        if matched_page:
            row.cells[1].text = matched_page
            p_c1 = row.cells[1].paragraphs[0]
            p_c1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for r in p_c1.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(11)

    doc.save(doc_path)
    print("Updated Table of Contents, Figures, and Tables successfully!", flush=True)

    # Quick final check
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        doc_w = word.Documents.Open(doc_path, ConfirmConversions=False, ReadOnly=False, AddToRecentFiles=False)
        final_pages = doc_w.ComputeStatistics(2)
        print(f"Final Verified Page Count: {final_pages}", flush=True)
        doc_w.Save()
        doc_w.Close(False)
    finally:
        word.Quit()

if __name__ == "__main__":
    update_toc()
