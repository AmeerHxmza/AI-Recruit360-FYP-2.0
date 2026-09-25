import docx

tpl = docx.Document('Project Report template-3.0.docx')
print('Sections:', len(tpl.sections))
for i, s in enumerate(tpl.sections):
    hdr = s.header
    ftr = s.footer
    hdr_txt = ' | '.join(p.text.strip() for p in hdr.paragraphs if p.text.strip())
    ftr_txt = ' | '.join(p.text.strip() for p in ftr.paragraphs if p.text.strip())
    print(f'Sec {i}: hdr="{hdr_txt}" ftr="{ftr_txt}"')
