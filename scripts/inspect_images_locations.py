import docx

for doc_name in ['FYP-THESIS.docx', 'FYP-Thesis-AIRecruit360.docx']:
    doc = docx.Document(doc_name)
    print(f"\n=== {doc_name} ===")
    print(f"Total Paragraphs: {len(doc.paragraphs)}, Tables: {len(doc.tables)}")
    drawings = []
    for i, p in enumerate(doc.paragraphs):
        for r_idx, r in enumerate(p.runs):
            d = r._r.xpath('.//w:drawing')
            if d:
                drawings.append((i, p.text[:50], doc.paragraphs[i+1].text[:50] if i+1 < len(doc.paragraphs) else ''))
    print(f"Total drawings found: {len(drawings)}")
    for item in drawings[:10]:
        print(f"  P{item[0]}: txt='{item[1]}' -> next='{item[2]}'")
