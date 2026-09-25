import docx

doc = docx.Document('FYP-Thesis-AIRecruit360.docx')
for t_idx, table in enumerate(doc.tables):
    aligns = set()
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                aligns.add(p.alignment)
    print(f"Table {t_idx}: alignments={aligns}")
