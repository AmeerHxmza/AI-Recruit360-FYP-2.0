import io
import pytest
from app.services.cv.extractor import extract_text_from_bytes
from app.core.exceptions import DocumentExtractionError

def test_pdf_text_resume():
    import fitz
    with fitz.open() as doc:
        doc.new_page().insert_text((72,72),'Built Python APIs with PostgreSQL')
        content=doc.tobytes()
    assert 'Python APIs' in extract_text_from_bytes(content,'resume.pdf','application/pdf')

def test_table_based_docx_resume():
    from docx import Document
    doc=Document();table=doc.add_table(rows=1,cols=2)
    table.cell(0,0).text='Experience';table.cell(0,1).text='Built Python APIs'
    data=io.BytesIO();doc.save(data)
    assert 'Built Python APIs' in extract_text_from_bytes(data.getvalue(),'resume.docx')

def test_empty_pdf_does_not_invent_evidence():
    import fitz
    with fitz.open() as doc:
        doc.new_page();content=doc.tobytes()
    with pytest.raises(DocumentExtractionError,match='OCR'):
        extract_text_from_bytes(content,'scan.pdf','application/pdf')
